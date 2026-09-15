import csv
from tkinter.font import Font
from urllib import request, response

from django.shortcuts import redirect, render
from django.http import HttpResponse

from django.shortcuts import render
from django.shortcuts import redirect
from django.shortcuts import get_object_or_404


from .models import Student, Attendance
from datetime import datetime, time

import io
import qrcode
from reportlab.lib.utils import ImageReader
from openpyxl import Workbook
from collections import defaultdict
from django.contrib.auth.decorators import login_required


@login_required
def dashboard(request):
    
    from datetime import date

    selected_date = request.GET.get("date")

    students = Student.objects.all()

    grouped_students = defaultdict(list)

    for student in students:
       
        latest_entry = Attendance.objects.filter(
            student=student
        ).last()
 
        if selected_date:
            last_entry = Attendance.objects.filter(
                student=student,
                date=selected_date
            ).last()
        else:
            last_entry = Attendance.objects.filter(
                student=student,
                date=date.today()
        ).last()

            latest_entry = Attendance.objects.filter(
                student=student
            ).last()


        if last_entry and last_entry.check_out is None:
            status = "✅ da"
        else:
            status = "❌ weg"

        if last_entry:
            check_in = last_entry.check_in
            check_out = last_entry.check_out
    
            if check_out is None:
                status = "✅"   # ist noch da
            else:
                status = "❌" 

        else:
            check_in = None
            check_out = None
            status = "❌"
        print(student.name, "| Tagesnotiz:", student.dashboard_note)


        grouped_students[student.student_class].append({
            "id": student.id,
            "student_id": student.student_id,
            "name": student.name,
            "class": student.student_class,
            "check_in": check_in.strftime("%H:%M") if check_in else "-",
            "check_out": check_out.strftime("%H:%M") if check_out else "-",
            "status": status,
            "attendance_id": latest_entry.id if latest_entry else None,
            "note": student.note,
            "dashboard_note": student.dashboard_note, 
        })
   
    
    
    return render(request, "dashboard.html", {
        "grouped_students": dict(grouped_students)
    })

from datetime import date

@login_required
def monthly_overview(request):
    today = date.today()

    students = Student.objects.all()

    results = []

    for student in students:
        entries = Attendance.objects.filter(
            student=student,
            date__month=today.month,
            date__year=today.year
        )

        total_cost = 0

        for entry in entries:
            if entry.check_in and entry.check_out:
                in_min = entry.check_in.hour * 60 + entry.check_in.minute
                out_min = entry.check_out.hour * 60 + entry.check_out.minute

                extra = 0

                if in_min < 450:   # vor 07:30
                    extra += 450 - in_min

                if out_min > 960:  # nach 16:00
                    extra += out_min - 960

                if extra > 0:
                    hours = (extra + 59) // 60
                    total_cost += hours * 6

        results.append({
            "student_id": student.student_id,
            "name": student.name,
            "class": student.student_class,
            "cost": total_cost
        })

    return render(request, "monthly.html", {
        "results": results
    })

 
from collections import defaultdict

def attendance_view(request):
    students = Student.objects.all()

    grouped_students = defaultdict(list)

    for student in students:
        grouped_students[student.class_name].append(student)

    return render(request, "attendance.html", {
        "grouped_students": grouped_students
    })
   

    

def generate_qr(request, student_id):
    img = qrcode.make(str(student_id))

    buffer = io.BytesIO()
    img.save(buffer, format='PNG')
    buffer.seek(0)

    return HttpResponse(buffer, content_type='image/png')

def import_students(request):
    import csv

    with open('schueler.csv', newline='', encoding='utf-8') as file:
        reader = csv.DictReader(file)

        for row in reader:
            Student.objects.get_or_create(
                name=list(row.values())[0].strip(),
                student_class=list(row.values())[1].strip()
            )

    return HttpResponse("Import erfolgreich ✅")

def scanner_page(request):
    return render(request, "scan.html")


# ✅ STATUS
def get_status(request):
    return HttpResponse("Status-Seite")


# ✅ SCAN
def scan_student(request, student_id):
    student = Student.objects.get(id=student_id)

    entry = Attendance.objects.filter(
        student=student,
        check_out__isnull=True
    ).first()

    if entry:
        entry.check_out = datetime.now().time()
        entry.save()
        
        return render(request, "scan_result.html", {
            "message": f"👋 Komm gut nach Hause, {student.name}!",
            "color": "#f7a600"
        })

    else:
        Attendance.objects.create(
            student=student,
            check_in=datetime.now().time()
        )
        
        return render(request, "scan_result.html", {
            "message": f"🌞 Guten Morgen, {student.name}! Schön, dass du da bist.",
            "color": "#84bd00"
        })



# ✅ QR EXPORT

def export_qr(request):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="qr_codes.pdf"'

    c = canvas.Canvas(response, pagesize=A4)
    width, height = A4

    card_width = 250
    card_height = 170

    margin_x = 50
    margin_y = 50

    gap_x = 40
    gap_y = 40

    x_start = margin_x
    y_start = height - margin_y - card_height

    x = x_start
    y = y_start

    count = 0

    for student in Student.objects.all():

        if y < margin_y:
            c.showPage()
            x = x_start
            y = height - margin_y - card_height

        img = qrcode.make(str(student.id))
        buffer = io.BytesIO()
        img.save(buffer, format="PNG")
        buffer.seek(0)

        image = ImageReader(buffer)

        c.rect(x, y, card_width, card_height)

        center_x = x + card_width / 2
        name_parts = student.name.split()

        if len(name_parts) >= 2:
            display_name = f"{name_parts[0]} {name_parts[-1][0]}."
        else:
            display_name = name_parts[0]
        name = display_name

        if len(name) > 30:
            part1 = name[:30]
            part2 = name[30:]
            c.drawCentredString(center_x, y + 135, part1)
            c.drawCentredString(center_x, y + 120, part2)
        else:
            c.drawCentredString(center_x, y + 125, name)

        qr_size = 70
        qr_x = center_x - qr_size / 2
        qr_y = y + 50

        c.drawImage(image, qr_x, qr_y, width=qr_size, height=qr_size)

        c.drawCentredString(center_x, y + 25, "Klasse " + student.student_class)

        if count % 2 == 0:
            x = x + card_width + gap_x
        else:
            x = x_start
            y = y - card_height - gap_y

        count += 1

    c.save()
    return response

    
@login_required
def export_excel(request):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Anwesenheit"

    ws.append(['ID', 'Name', 'Klasse', 'Kommen', 'Gehen', 'Kosten (€)'])
    from openpyxl.styles import Font
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles import Border, Font, PatternFill
    from openpyxl.styles import Font, PatternFill, Side
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill(
        start_color="84BD00",
        end_color="84BD00",
        fill_type="solid"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Kopfzeile formatieren
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    students = Student.objects.all()
    print("Anzahl Schüler:", students.count())

    for student in students:
        print("Exportiere:", student.name)

    for student in students:
        from datetime import date

        last_entry = Attendance.objects.filter(
            student=student,
            date=date.today()
        ).last()


        if last_entry:
            check_in = last_entry.check_in
            check_out = last_entry.check_out

        else:
            check_in = None
            check_out = None

        cost = 0
    
        ws.append([
            student.student_id,
            student.name,
            student.student_class,
            check_in.strftime("%H:%M") if isinstance(check_in, time) else "-",
            check_out.strftime("%H:%M") if isinstance(check_out, time) else "-",
            cost
        ])
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = (
        'attachment; filename="anwesenheit.xlsx"'
    )
    ws.auto_filter.ref = ws.dimensions
    
    
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass

        ws.column_dimensions[column_letter].width = max_length + 2   
        ws.freeze_panes = "A2" 
    wb.save(response)    
    return response
@login_required   
def monthly_report(request):
    import csv
    from datetime import date

    today = date.today()

    selected_class = request.GET.get("class")
    month_string = request.GET.get("month")

    if month_string:
        year, month = month_string.split("-")
        year = int(year)
        month = int(month)
    else:
        year = today.year
        month = today.month

    response = HttpResponse(
        content_type='text/csv; charset=utf-8'
    )
    response.write('\ufeff')

    response['Content-Disposition'] = (
        'attachment; filename="monat.csv"'
    )

    writer = csv.writer(response, delimiter=';')

    month_names = {
        1: "Januar",
        2: "Februar",
        3: "März",
        4: "April",
        5: "Mai",
        6: "Juni",
        7: "Juli",
        8: "August",
        9: "September",
        10: "Oktober",
        11: "November",
        12: "Dezember"
    }

    writer.writerow([])
    writer.writerow([
        f"Monatsabrechnung {month_names[month]} {year}"
    ])
    writer.writerow([])

    writer.writerow([
        'ID',
        'Name',
        'Klasse',
        'Monatskosten (€)',
        'Kosten entstanden am'
    ])

    students = Student.objects.all()


    if selected_class:
        students = students.filter(
            student_class=selected_class
        )

    print("Klasse:", selected_class)
    print("Anzahl Schüler:", students.count())

    for student in students:
        entries = Attendance.objects.filter(
            student=student,
            date__month=month,
            date__year=year
        )
        total_cost = 0
        cost_dates = []
        for entry in entries:
            if entry.check_in and entry.check_out:

                in_min = entry.check_in.hour * 60 + entry.check_in.minute
                out_min = entry.check_out.hour * 60 + entry.check_out.minute

                extra = 0

                if in_min < 450:  # vor 7:30
                    extra += max(0, 450 - in_min)

                if out_min > 960:  # nach 16:00
                    extra += out_min - 960

                if extra > 0:
                    hours = (extra + 59) // 60
                    total_cost += hours * 6
                    cost_dates.append(
                        entry.date.strftime("%d.%m")
                    )

        cost_dates_text = ", ".join(cost_dates)
        writer.writerow([
            student.student_id,
            student.name,
            student.student_class,
            total_cost,
            cost_dates_text
        ])


    return response


@login_required
def edit_attendance(request, attendance_id):

    attendance = get_object_or_404(
        Attendance,
        id=attendance_id
    )

    if request.method == "POST":
        check_in = request.POST.get("check_in")
        check_out = request.POST.get("check_out")
        dashboard_note = request.POST.get("dashboard_note")

        attendance.check_in = check_in if check_in else None
        attendance.check_out = check_out if check_out else None
        attendance.student.dashboard_note = dashboard_note
        

        attendance.save()
        attendance.student.save()
        return redirect("dashboard")

    return render(
        request,
        "edit_attendance.html",
        {"attendance": attendance}
    )
@login_required
def new_attendance(request, student_id):

    student = get_object_or_404(
        Student,
        id=student_id
    )

    if request.method == "POST":

        check_in = request.POST.get("check_in")
        check_out = request.POST.get("check_out")
        

        Attendance.objects.create(
            student=student,
            date=date.today(),
            check_in=check_in,
            check_out=check_out,
            
        )

        return redirect("dashboard")

    return render(
        request,
        "new_attendance.html",
        {"student": student}
    )